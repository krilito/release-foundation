from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
BASE = ROOT / "docs" / "tables" / "chitosan"
OUT = BASE / "jianbiao.xlsx"
SAMPLE = BASE / "sample.csv"
RELEASE = BASE / "release.csv"
VALIDATION = BASE / "validation_report.txt"

HEADER_FILL = PatternFill("solid", fgColor="4F4F4F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUB_FILL = PatternFill("solid", fgColor="D9D9D9")
WRAP = Alignment(vertical="top", wrap_text=True)
def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def to_float(text: str | None) -> float | None:
    if text is None:
        return None
    s = str(text).strip()
    if not s:
        return None
    return float(s)


def style(ws, freeze: str = "A2") -> None:
    from openpyxl.utils import get_column_letter

    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = WRAP
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width, 10), 28)


def build():
    sample_rows = read_csv(SAMPLE)
    release_rows = read_csv(RELEASE)

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("说明")
    lines = [
        "这是简版，给人直接看和直接填。",
        "原则：先看懂，再分析。",
        "",
        "工作表：",
        "1. 样本：每组是什么、缺什么。",
        "2. 释药：按时间展开的 DEX 数据。",
        "3. 参考：论文参考值，单独放。",
        "4. 问题：现在最该回查的点。",
        "",
        "已锁定映射：",
        "组1=F2，组2=F3，组3=F5，组4=F6，组5=F1，组6=F4。",
        "",
        "注意：",
        "1. HA 这里按 20-40万 / 40-80万 Da 理解，对应 0.2-0.4 / 0.4-0.8 MDa。",
        "2. 师兄论文里的 80-150万 Da 是 HA，不是 CS。",
        "3. 现在很多 Qt>100%，说明原始吸光度/分母/标曲还要回查。",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(i, 1, line)
    ws["A1"].fill = SUB_FILL
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    ws.freeze_panes = "A1"

    ws = wb.create_sheet("样本")
    sample_header = ["F号", "组号", "药", "HA(MDa)", "PVP(kDa)", "DEX(mg)", "GCV(mg)", "粒径", "PDI", "Zeta", "状态", "备注"]
    ws.append(sample_header)
    for row in sample_rows:
        group = row["实验原始组号"]
        if group.startswith("参考-"):
            continue
        ha = ""
        if row["HA名义分子量_MDa"]:
            raw = row["HA名义分子量_MDa"].strip()
            if raw == "20-40":
                ha = "0.2-0.4"
            elif raw == "40-80":
                ha = "0.4-0.8"
            else:
                ha = raw
        pvp = row["PVP名义分子量_kDa"] or "无"
        note = row["异常备注"] or ""
        status = "待补"
        ws.append([
            row["锁定配方编号F1_F6"],
            group,
            row["药物"],
            ha,
            pvp,
            to_float(row["DEX理论投料量_mg"]),
            to_float(row["GCV理论投料量_mg"]),
            to_float(row["粒径_nm"]),
            to_float(row["PDI"]),
            to_float(row["Zeta电位_mV"]),
            status,
            note,
        ])
    style(ws, "A2")

    ws = wb.create_sheet("释药")
    release_header = ["F号", "组号", "药", "时间(h)", "释放(%)", "原吸", "空白", "倍数", "状态", "备注"]
    ws.append(release_header)
    for row in release_rows:
        pct = to_float(row["累计释放率_pct"])
        status = "异常" if pct is not None and pct > 100 else "正常"
        ws.append([
            row["锁定配方编号F1_F6"],
            row["实验原始组号"],
            row["药物"],
            to_float(row["时间_h"]),
            pct,
            to_float(row["原始吸光度"]),
            to_float(row["空白吸光度"]),
            to_float(row["稀释倍数"]),
            status,
            row["备注"],
        ])
    style(ws, "A2")

    ws = wb.create_sheet("参考")
    ref_header = ["对象", "载药率", "包封率", "粒径", "Zeta", "备注"]
    ws.append(ref_header)
    for row in sample_rows:
        group = row["实验原始组号"]
        if not group.startswith("参考-"):
            continue
        ws.append([
            group.replace("参考-", ""),
            row["DEX实际装载量"],
            row["DEX包封率_pct"],
            to_float(row["粒径_nm"]),
            to_float(row["Zeta电位_mV"]),
            row["异常备注"],
        ])
    style(ws, "A2")

    ws = wb.create_sheet("问题")
    problem_header = ["级别", "问题", "处理"]
    ws.append(problem_header)
    problems = [
        ("高", "很多 DEX 累计释放率 > 100%", "先回查原始吸光度、空白、稀释倍数、Qt分母"),
        ("高", "现在 release 表只有 mean，没有3次重复原始值", "补每个时间点的独立重复"),
        ("中", "粒径 / PDI / Zeta 还没补到6组", "每组分别测，不要借师兄单一配方"),
        ("中", "HA 单位原始写法容易误读", "统一按 0.2-0.4 / 0.4-0.8 MDa 看"),
        ("低", "GCV 释放数据还没进来", "后续补 GCV 支线"),
    ]
    for row in problems:
        ws.append(row)

    if VALIDATION.exists():
        ws.append(("", "", ""))
        ws.append(("附", "数据库版校验摘要", "只摘最关键的问题"))
        text = VALIDATION.read_text(encoding="utf-8")
        count = 0
        for line in text.splitlines():
            if "[warning]" in line and "Percent exceeds 100." in line:
                ws.append(("warning", "存在 >100% 释放点", line))
                count += 1
                if count >= 5:
                    break
    style(ws, "A2")

    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
