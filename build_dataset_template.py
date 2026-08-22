from __future__ import annotations

import csv
import re
import shutil
from collections import defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent
INPUT_DIR = ROOT / "docs" / "tables" / "chitosan"
OUTPUT_XLSX = INPUT_DIR / "chitosan_hydrogel_dataset_template_v2.xlsx"
VALIDATION_TXT = INPUT_DIR / "validation_report.txt"
CHANGELOG_MD = INPUT_DIR / "CHANGELOG.md"
BACKUP_DIR = INPUT_DIR / "backup"

SOURCE_WORKBOOK = INPUT_DIR / "释放曲线总表.xlsx"
SOURCE_SAMPLE_XLSX = INPUT_DIR / "sample.xlsx"
SOURCE_RELEASE_XLSX = INPUT_DIR / "release.xlsx"
SOURCE_SAMPLE_CSV = INPUT_DIR / "sample.csv"
SOURCE_RELEASE_CSV = INPUT_DIR / "release.csv"

HEADER_FILL = PatternFill("solid", fgColor="4F4F4F")
SUB_FILL = PatternFill("solid", fgColor="D9D9D9")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(vertical="top", wrap_text=True)
TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)

ENUMS = {
    "record_type": ["experiment", "literature_reference", "pilot", "validation"],
    "qc_status": ["pass", "review", "exclude", "not_evaluated"],
    "data_quality": ["directly_reported", "digitized_from_figure", "calculated", "uncertain"],
}

SHEETS = {
    "Formulation": [
        "formulation_id",
        "record_type",
        "formulation_name",
        "sample_group",
        "cs_mw_min_kda",
        "cs_mw_max_kda",
        "cs_deacetylation_degree_pct",
        "cs_concentration_mg_ml",
        "cs_supplier",
        "cs_catalog_number",
        "cs_batch_number",
        "tpp_mw_kda",
        "tpp_concentration_mg_ml",
        "ha_mw_grade",
        "ha_mw_min_mda",
        "ha_mw_max_mda",
        "ha_concentration_mg_ml",
        "ha_supplier",
        "ha_catalog_number",
        "pvp_grade",
        "pvp_nominal_mw_kda",
        "pvp_concentration_mg_ml",
        "dex_feed_mass_mg",
        "gcv_feed_mass_mg",
        "hydrogel_type",
        "formulation_notes",
        "source_id",
    ],
    "Process": [
        "batch_id",
        "formulation_id",
        "preparation_date",
        "operator",
        "replicate_id",
        "biological_or_independent_replicate",
        "cs_solution_ph",
        "ha_solution_ph",
        "tpp_solution_ph",
        "addition_order",
        "addition_rate_ml_min",
        "stirring_speed_rpm",
        "stirring_time_min",
        "reaction_temperature_c",
        "centrifuge_speed_rcf",
        "centrifuge_speed_rpm",
        "centrifuge_time_min",
        "washing_cycles",
        "freeze_drying",
        "storage_temperature_c",
        "process_notes",
    ],
    "Particle_Result": [
        "particle_result_id",
        "batch_id",
        "measurement_date",
        "technical_replicate_id",
        "dispersion_medium",
        "measurement_temperature_c",
        "instrument_model",
        "particle_size_mean_nm",
        "particle_size_sd_nm",
        "pdi_mean",
        "pdi_sd",
        "zeta_potential_mean_mv",
        "zeta_potential_sd_mv",
        "dex_recovered_mass_mean_mg",
        "dex_recovered_mass_sd_mg",
        "dex_drug_loading_mean_pct",
        "dex_drug_loading_sd_pct",
        "dex_encapsulation_efficiency_mean_pct",
        "dex_encapsulation_efficiency_sd_pct",
        "gcv_recovered_mass_mean_mg",
        "gcv_recovered_mass_sd_mg",
        "gcv_drug_loading_mean_pct",
        "gcv_drug_loading_sd_pct",
        "gcv_encapsulation_efficiency_mean_pct",
        "gcv_encapsulation_efficiency_sd_pct",
        "result_notes",
        "qc_status",
    ],
    "Hydrogel_Result": [
        "hydrogel_result_id",
        "batch_id",
        "measurement_date",
        "technical_replicate_id",
        "gelation_time_s",
        "injectability_force_n",
        "storage_modulus_pa",
        "loss_modulus_pa",
        "frequency_hz",
        "strain_pct",
        "swelling_time_h",
        "swelling_ratio_mean_pct",
        "swelling_ratio_sd_pct",
        "degradation_time_day",
        "degradation_ratio_mean_pct",
        "degradation_ratio_sd_pct",
        "dex_release_time_h",
        "dex_cumulative_release_mean_pct",
        "dex_cumulative_release_sd_pct",
        "gcv_release_time_h",
        "gcv_cumulative_release_mean_pct",
        "gcv_cumulative_release_sd_pct",
        "cell_viability_mean_pct",
        "cell_viability_sd_pct",
        "result_notes",
        "qc_status",
    ],
    "Literature_Reference": [
        "source_id",
        "reference_type",
        "authors",
        "title",
        "year",
        "institution_or_journal",
        "doi",
        "document_page",
        "pdf_page",
        "chapter",
        "section",
        "table_number",
        "figure_number",
        "material_name",
        "parameter_name",
        "value_mean",
        "value_sd",
        "value_min",
        "value_max",
        "unit",
        "extraction_method",
        "data_quality",
        "notes",
    ],
    "Data_Dictionary": [
        "sheet_name",
        "field_name",
        "chinese_name",
        "definition",
        "data_type",
        "unit",
        "allowed_values",
        "required",
        "missing_value_rule",
        "example",
        "remarks",
    ],
    "Validation_Report": [
        "check_id",
        "sheet_name",
        "row_number",
        "field_name",
        "severity",
        "problem",
        "suggested_action",
    ],
}

NUMERIC_FIELDS = {
    "Formulation": {
        "cs_mw_min_kda", "cs_mw_max_kda", "cs_deacetylation_degree_pct", "cs_concentration_mg_ml",
        "tpp_mw_kda", "tpp_concentration_mg_ml", "ha_mw_min_mda", "ha_mw_max_mda",
        "ha_concentration_mg_ml", "pvp_nominal_mw_kda", "pvp_concentration_mg_ml",
        "dex_feed_mass_mg", "gcv_feed_mass_mg",
    },
    "Process": {
        "cs_solution_ph", "ha_solution_ph", "tpp_solution_ph", "addition_rate_ml_min",
        "stirring_speed_rpm", "stirring_time_min", "reaction_temperature_c",
        "centrifuge_speed_rcf", "centrifuge_speed_rpm", "centrifuge_time_min",
        "washing_cycles", "storage_temperature_c",
    },
    "Particle_Result": {
        "measurement_temperature_c", "particle_size_mean_nm", "particle_size_sd_nm",
        "pdi_mean", "pdi_sd", "zeta_potential_mean_mv", "zeta_potential_sd_mv",
        "dex_recovered_mass_mean_mg", "dex_recovered_mass_sd_mg",
        "dex_drug_loading_mean_pct", "dex_drug_loading_sd_pct",
        "dex_encapsulation_efficiency_mean_pct", "dex_encapsulation_efficiency_sd_pct",
        "gcv_recovered_mass_mean_mg", "gcv_recovered_mass_sd_mg",
        "gcv_drug_loading_mean_pct", "gcv_drug_loading_sd_pct",
        "gcv_encapsulation_efficiency_mean_pct", "gcv_encapsulation_efficiency_sd_pct",
    },
    "Hydrogel_Result": {
        "gelation_time_s", "injectability_force_n", "storage_modulus_pa", "loss_modulus_pa",
        "frequency_hz", "strain_pct", "swelling_time_h", "swelling_ratio_mean_pct",
        "swelling_ratio_sd_pct", "degradation_time_day", "degradation_ratio_mean_pct",
        "degradation_ratio_sd_pct", "dex_release_time_h", "dex_cumulative_release_mean_pct",
        "dex_cumulative_release_sd_pct", "gcv_release_time_h", "gcv_cumulative_release_mean_pct",
        "gcv_cumulative_release_sd_pct", "cell_viability_mean_pct", "cell_viability_sd_pct",
    },
    "Literature_Reference": {"year", "document_page", "pdf_page", "value_mean", "value_sd", "value_min", "value_max"},
}

ID_FIELDS = {
    "Formulation": {"formulation_id", "source_id"},
    "Process": {"batch_id", "formulation_id", "replicate_id"},
    "Particle_Result": {"particle_result_id", "batch_id", "technical_replicate_id"},
    "Hydrogel_Result": {"hydrogel_result_id", "batch_id", "technical_replicate_id"},
    "Literature_Reference": {"source_id"},
    "Validation_Report": {"check_id"},
}

REQUIRED_FIELDS = {
    "Formulation": {"formulation_id", "record_type", "formulation_name", "sample_group", "source_id"},
    "Process": {"batch_id", "formulation_id"},
    "Particle_Result": {"particle_result_id", "batch_id", "qc_status"},
    "Hydrogel_Result": {"hydrogel_result_id", "batch_id", "qc_status"},
    "Literature_Reference": {"source_id", "reference_type", "material_name", "parameter_name", "unit", "data_quality"},
}

CHINESE_NAMES = {
    "formulation_id": "配方ID",
    "record_type": "记录类型",
    "formulation_name": "配方名称",
    "sample_group": "样品组号",
    "cs_mw_min_kda": "壳聚糖分子量下限",
    "cs_mw_max_kda": "壳聚糖分子量上限",
    "cs_deacetylation_degree_pct": "壳聚糖脱乙酰度",
    "cs_concentration_mg_ml": "壳聚糖浓度",
    "cs_supplier": "壳聚糖供应商",
    "cs_catalog_number": "壳聚糖货号",
    "cs_batch_number": "壳聚糖批号",
    "tpp_mw_kda": "TPP分子量",
    "tpp_concentration_mg_ml": "TPP浓度",
    "ha_mw_grade": "HA分子量等级",
    "ha_mw_min_mda": "HA分子量下限",
    "ha_mw_max_mda": "HA分子量上限",
    "ha_concentration_mg_ml": "HA浓度",
    "ha_supplier": "HA供应商",
    "ha_catalog_number": "HA货号",
    "pvp_grade": "PVP牌号",
    "pvp_nominal_mw_kda": "PVP标称分子量",
    "pvp_concentration_mg_ml": "PVP浓度",
    "dex_feed_mass_mg": "DEX理论投料量",
    "gcv_feed_mass_mg": "GCV理论投料量",
    "hydrogel_type": "水凝胶类型",
    "formulation_notes": "配方备注",
    "source_id": "来源ID",
    "batch_id": "批次ID",
    "preparation_date": "制备日期",
    "operator": "操作人",
    "replicate_id": "独立重复编号",
    "biological_or_independent_replicate": "独立重复说明",
    "cs_solution_ph": "CS溶液pH",
    "ha_solution_ph": "HA溶液pH",
    "tpp_solution_ph": "TPP溶液pH",
    "addition_order": "加料顺序",
    "addition_rate_ml_min": "加料速度",
    "stirring_speed_rpm": "搅拌转速",
    "stirring_time_min": "搅拌时间",
    "reaction_temperature_c": "反应温度",
    "centrifuge_speed_rcf": "离心力",
    "centrifuge_speed_rpm": "离心转速",
    "centrifuge_time_min": "离心时间",
    "washing_cycles": "洗涤次数",
    "freeze_drying": "冻干",
    "storage_temperature_c": "储存温度",
    "process_notes": "工艺备注",
    "particle_result_id": "颗粒结果ID",
    "measurement_date": "测量日期",
    "technical_replicate_id": "技术重复编号",
    "dispersion_medium": "分散介质",
    "measurement_temperature_c": "测量温度",
    "instrument_model": "仪器型号",
    "particle_size_mean_nm": "粒径均值",
    "particle_size_sd_nm": "粒径标准差",
    "pdi_mean": "PDI均值",
    "pdi_sd": "PDI标准差",
    "zeta_potential_mean_mv": "Zeta电位均值",
    "zeta_potential_sd_mv": "Zeta电位标准差",
    "dex_recovered_mass_mean_mg": "DEX回收质量均值",
    "dex_recovered_mass_sd_mg": "DEX回收质量标准差",
    "dex_drug_loading_mean_pct": "DEX载药率均值",
    "dex_drug_loading_sd_pct": "DEX载药率标准差",
    "dex_encapsulation_efficiency_mean_pct": "DEX包封率均值",
    "dex_encapsulation_efficiency_sd_pct": "DEX包封率标准差",
    "gcv_recovered_mass_mean_mg": "GCV回收质量均值",
    "gcv_recovered_mass_sd_mg": "GCV回收质量标准差",
    "gcv_drug_loading_mean_pct": "GCV载药率均值",
    "gcv_drug_loading_sd_pct": "GCV载药率标准差",
    "gcv_encapsulation_efficiency_mean_pct": "GCV包封率均值",
    "gcv_encapsulation_efficiency_sd_pct": "GCV包封率标准差",
    "result_notes": "结果备注",
    "qc_status": "质控状态",
    "hydrogel_result_id": "水凝胶结果ID",
    "gelation_time_s": "凝胶时间",
    "injectability_force_n": "注射力",
    "storage_modulus_pa": "储能模量",
    "loss_modulus_pa": "损耗模量",
    "frequency_hz": "频率",
    "strain_pct": "应变",
    "swelling_time_h": "溶胀时间",
    "swelling_ratio_mean_pct": "溶胀比均值",
    "swelling_ratio_sd_pct": "溶胀比标准差",
    "degradation_time_day": "降解时间",
    "degradation_ratio_mean_pct": "降解比例均值",
    "degradation_ratio_sd_pct": "降解比例标准差",
    "dex_release_time_h": "DEX释放时间",
    "dex_cumulative_release_mean_pct": "DEX累计释放均值",
    "dex_cumulative_release_sd_pct": "DEX累计释放标准差",
    "gcv_release_time_h": "GCV释放时间",
    "gcv_cumulative_release_mean_pct": "GCV累计释放均值",
    "gcv_cumulative_release_sd_pct": "GCV累计释放标准差",
    "cell_viability_mean_pct": "细胞活性均值",
    "cell_viability_sd_pct": "细胞活性标准差",
    "reference_type": "参考类型",
    "authors": "作者",
    "title": "标题",
    "year": "年份",
    "institution_or_journal": "机构或期刊",
    "doi": "DOI",
    "document_page": "文档页码",
    "pdf_page": "PDF页码",
    "chapter": "章节",
    "section": "小节",
    "table_number": "表号",
    "figure_number": "图号",
    "material_name": "材料名称",
    "parameter_name": "参数名称",
    "value_mean": "均值",
    "value_sd": "标准差",
    "value_min": "最小值",
    "value_max": "最大值",
    "unit": "单位",
    "extraction_method": "提取方式",
    "data_quality": "数据质量",
    "sheet_name": "工作表名",
    "field_name": "字段名",
    "definition": "定义",
    "data_type": "数据类型",
    "allowed_values": "允许值",
    "required": "是否必填",
    "missing_value_rule": "缺失值规则",
    "example": "示例",
    "remarks": "备注",
    "check_id": "检查ID",
    "row_number": "行号",
    "severity": "严重程度",
    "problem": "问题",
    "suggested_action": "建议处理",
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def to_float(value: str | None) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "")
    return float(text)


def parse_mean_sd(text: str | None) -> tuple[float | None, float | None]:
    if text is None:
        return (None, None)
    s = str(text).strip()
    if not s:
        return (None, None)
    s = s.replace("（", "(").replace("）", ")")
    s = s.replace("%", "")
    match = re.search(r"([-+]?\d+(?:\.\d+)?)\s*±\s*([-+]?\d+(?:\.\d+)?)", s)
    if match:
        return (float(match.group(1)), float(match.group(2)))
    if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", s):
        return (float(s), None)
    return (None, None)


def parse_ha_range_to_mda(text: str | None) -> tuple[float | None, float | None]:
    if text is None:
        return (None, None)
    s = str(text).strip()
    if not s:
        return (None, None)
    numbers = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", s)]
    if len(numbers) < 2:
        return (None, None)
    low, high = numbers[0], numbers[1]
    if high >= 10:
        return (round(low / 100.0, 4), round(high / 100.0, 4))
    return (low, high)


def formulation_name(sample_group: str, ha_min: float | None, ha_max: float | None, pvp: float | None) -> str:
    ha_label = "unknown_HA" if ha_min is None or ha_max is None else f"HA_{ha_min:g}_{ha_max:g}_MDa"
    pvp_label = "no_PVP" if pvp in (None, 0) else f"PVP_{int(pvp)}_kDa"
    return f"{sample_group}_{ha_label}_{pvp_label}"


def backup_sources() -> list[Path]:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backups: list[Path] = []
    for src in [SOURCE_WORKBOOK, SOURCE_SAMPLE_XLSX, SOURCE_RELEASE_XLSX, SOURCE_SAMPLE_CSV, SOURCE_RELEASE_CSV]:
        if src.exists():
            dst = BACKUP_DIR / f"{src.stem}_backup_{stamp}{src.suffix}"
            shutil.copy2(src, dst)
            backups.append(dst)
    return backups


def build_formulation_rows(sample_rows: list[dict[str, str]]) -> tuple[list[dict[str, object]], dict[str, str]]:
    rows: list[dict[str, object]] = []
    batch_map: dict[str, str] = {}
    for row in sample_rows:
        group = row["实验原始组号"]
        locked = row["锁定配方编号F1_F6"].strip() if row["锁定配方编号F1_F6"] else ""
        is_reference = group.startswith("参考-")
        source_id = "SRC-ZYF-THESIS-2024" if is_reference else "SRC-QDU-CHITOSAN-LOCAL"
        if is_reference:
            formulation_id = {
                "参考-DCNPs(裸核)": "REF-ZYF-01",
                "参考-HDCNPs(+HA)": "REF-ZYF-02",
                "参考-PHDCNPs(+HA+PVP)": "REF-ZYF-03",
            }.get(group, f"REF-{len(rows)+1:02d}")
            record_type = "literature_reference"
        else:
            formulation_id = f"EXP-{locked}"
            batch_map[locked] = f"{formulation_id}-B01"
            record_type = "experiment"

        ha_min, ha_max = parse_ha_range_to_mda(row["HA名义分子量_MDa"])
        pvp = to_float(row["PVP名义分子量_kDa"])
        rows.append(
            {
                "formulation_id": formulation_id,
                "record_type": record_type,
                "formulation_name": formulation_name(group, ha_min, ha_max, pvp),
                "sample_group": group,
                "cs_mw_min_kda": None,
                "cs_mw_max_kda": None,
                "cs_deacetylation_degree_pct": None,
                "cs_concentration_mg_ml": None,
                "cs_supplier": None,
                "cs_catalog_number": None,
                "cs_batch_number": None,
                "tpp_mw_kda": None,
                "tpp_concentration_mg_ml": None,
                "ha_mw_grade": row["HA分子量档"] or None,
                "ha_mw_min_mda": ha_min,
                "ha_mw_max_mda": ha_max,
                "ha_concentration_mg_ml": None,
                "ha_supplier": None,
                "ha_catalog_number": None,
                "pvp_grade": "K30" if pvp == 58 else None,
                "pvp_nominal_mw_kda": pvp,
                "pvp_concentration_mg_ml": None,
                "dex_feed_mass_mg": to_float(row["DEX理论投料量_mg"]),
                "gcv_feed_mass_mg": to_float(row["GCV理论投料量_mg"]),
                "hydrogel_type": None,
                "formulation_notes": (
                    (row["异常备注"] or "").strip() + "; cs_mw and cs_deacetylation not_reported"
                ).strip("; "),
                "source_id": source_id,
            }
        )
    return rows, batch_map


def build_process_rows(formulation_rows: list[dict[str, object]], batch_map: dict[str, str]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for f in formulation_rows:
        if f["record_type"] != "experiment":
            continue
        formulation_id = str(f["formulation_id"])
        locked = formulation_id.replace("EXP-", "")
        rows.append(
            {
                "batch_id": batch_map[locked],
                "formulation_id": formulation_id,
                "preparation_date": None,
                "operator": None,
                "replicate_id": None,
                "biological_or_independent_replicate": None,
                "cs_solution_ph": None,
                "ha_solution_ph": None,
                "tpp_solution_ph": None,
                "addition_order": None,
                "addition_rate_ml_min": None,
                "stirring_speed_rpm": None,
                "stirring_time_min": None,
                "reaction_temperature_c": None,
                "centrifuge_speed_rcf": None,
                "centrifuge_speed_rpm": None,
                "centrifuge_time_min": None,
                "washing_cycles": None,
                "freeze_drying": None,
                "storage_temperature_c": None,
                "process_notes": "not_reported",
            }
        )
    return rows


def build_particle_rows(formulation_rows: list[dict[str, object]], batch_map: dict[str, str]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    counter = 1
    for f in formulation_rows:
        if f["record_type"] != "experiment":
            continue
        locked = str(f["formulation_id"]).replace("EXP-", "")
        rows.append(
            {
                "particle_result_id": f"PR-{counter:03d}",
                "batch_id": batch_map[locked],
                "measurement_date": None,
                "technical_replicate_id": None,
                "dispersion_medium": None,
                "measurement_temperature_c": None,
                "instrument_model": None,
                "particle_size_mean_nm": None,
                "particle_size_sd_nm": None,
                "pdi_mean": None,
                "pdi_sd": None,
                "zeta_potential_mean_mv": None,
                "zeta_potential_sd_mv": None,
                "dex_recovered_mass_mean_mg": None,
                "dex_recovered_mass_sd_mg": None,
                "dex_drug_loading_mean_pct": None,
                "dex_drug_loading_sd_pct": None,
                "dex_encapsulation_efficiency_mean_pct": None,
                "dex_encapsulation_efficiency_sd_pct": None,
                "gcv_recovered_mass_mean_mg": None,
                "gcv_recovered_mass_sd_mg": None,
                "gcv_drug_loading_mean_pct": None,
                "gcv_drug_loading_sd_pct": None,
                "gcv_encapsulation_efficiency_mean_pct": None,
                "gcv_encapsulation_efficiency_sd_pct": None,
                "result_notes": "template row; result not_reported in current source",
                "qc_status": "not_evaluated",
            }
        )
        counter += 1
    return rows


def build_hydrogel_rows(release_rows: list[dict[str, str]], batch_map: dict[str, str]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for idx, row in enumerate(release_rows, start=1):
        locked = row["锁定配方编号F1_F6"]
        pct = to_float(row["累计释放率_pct"])
        qc = "review" if pct is not None and pct > 100 else "not_evaluated"
        rows.append(
            {
                "hydrogel_result_id": f"HR-{idx:04d}",
                "batch_id": batch_map.get(locked),
                "measurement_date": None,
                "technical_replicate_id": row["重复编号"] or None,
                "gelation_time_s": None,
                "injectability_force_n": None,
                "storage_modulus_pa": None,
                "loss_modulus_pa": None,
                "frequency_hz": None,
                "strain_pct": None,
                "swelling_time_h": None,
                "swelling_ratio_mean_pct": None,
                "swelling_ratio_sd_pct": None,
                "degradation_time_day": None,
                "degradation_ratio_mean_pct": None,
                "degradation_ratio_sd_pct": None,
                "dex_release_time_h": to_float(row["时间_h"]),
                "dex_cumulative_release_mean_pct": pct,
                "dex_cumulative_release_sd_pct": None,
                "gcv_release_time_h": None,
                "gcv_cumulative_release_mean_pct": None,
                "gcv_cumulative_release_sd_pct": None,
                "cell_viability_mean_pct": None,
                "cell_viability_sd_pct": None,
                "result_notes": row["备注"] or "mean release imported from release.csv",
                "qc_status": qc,
            }
        )
    return rows


def build_literature_rows(sample_rows: list[dict[str, str]]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = [
        {
            "source_id": "SRC-ZYF-THESIS-2024",
            "reference_type": "thesis",
            "authors": "Zhang Yongfei",
            "title": "序贯药物控释可注射复合水凝胶的制备与性能研究",
            "year": 2024,
            "institution_or_journal": "Qingdao University",
            "doi": None,
            "document_page": None,
            "pdf_page": None,
            "chapter": "第三章",
            "section": "第二体系",
            "table_number": None,
            "figure_number": None,
            "material_name": "hyaluronic_acid",
            "parameter_name": "molecular_weight",
            "value_mean": None,
            "value_sd": None,
            "value_min": 0.8,
            "value_max": 1.5,
            "unit": "MDa",
            "extraction_method": "manual_from_notes",
            "data_quality": "directly_reported",
            "notes": "学长论文中的80-150万Da对应HA，不是CS。",
        }
    ]
    name_map = {
        "参考-DCNPs(裸核)": "DCNPs",
        "参考-HDCNPs(+HA)": "HDCNPs",
        "参考-PHDCNPs(+HA+PVP)": "PHDCNPs",
    }
    for row in sample_rows:
        group = row["实验原始组号"]
        if not group.startswith("参考-"):
            continue
        material = name_map[group]
        dl_mean, dl_sd = parse_mean_sd(row["DEX实际装载量"])
        ee_mean, ee_sd = parse_mean_sd(row["DEX包封率_pct"])
        size = to_float(row["粒径_nm"])
        zeta = to_float(row["Zeta电位_mV"])
        for parameter_name, mean, sd, unit in [
            ("dex_drug_loading", dl_mean, dl_sd, "pct"),
            ("dex_encapsulation_efficiency", ee_mean, ee_sd, "pct"),
            ("particle_size", size, None, "nm"),
            ("zeta_potential", zeta, None, "mV"),
        ]:
            rows.append(
                {
                    "source_id": "SRC-ZYF-THESIS-2024",
                    "reference_type": "thesis",
                    "authors": "Zhang Yongfei",
                    "title": "序贯药物控释可注射复合水凝胶的制备与性能研究",
                    "year": 2024,
                    "institution_or_journal": "Qingdao University",
                    "doi": None,
                    "document_page": None,
                    "pdf_page": None,
                    "chapter": "第三章",
                    "section": "第二体系",
                    "table_number": None,
                    "figure_number": None,
                    "material_name": material,
                    "parameter_name": parameter_name,
                    "value_mean": mean,
                    "value_sd": sd,
                    "value_min": None,
                    "value_max": None,
                    "unit": unit,
                    "extraction_method": "manual_from_reference_rows",
                    "data_quality": "directly_reported",
                    "notes": row["颗粒或凝胶批次备注"] or row["异常备注"],
                }
            )
    return rows


def dict_row(sheet_name: str, field_name: str) -> dict[str, object]:
    allowed = ",".join(ENUMS.get(field_name, []))
    unit = ""
    if field_name.endswith("_pct"):
        unit = "pct"
    elif field_name.endswith("_mg"):
        unit = "mg"
    elif field_name.endswith("_mda"):
        unit = "MDa"
    elif field_name.endswith("_kda"):
        unit = "kDa"
    elif field_name.endswith("_nm"):
        unit = "nm"
    elif field_name.endswith("_mv"):
        unit = "mV"
    elif field_name.endswith("_rpm"):
        unit = "rpm"
    elif field_name.endswith("_rcf"):
        unit = "rcf"
    elif field_name.endswith("_c"):
        unit = "C"
    elif field_name.endswith("_ml") or field_name.endswith("_ml_min"):
        unit = "mL or mL/min"
    elif field_name.endswith("_min"):
        unit = "min"
    elif field_name.endswith("_day"):
        unit = "day"
    elif field_name.endswith("_h"):
        unit = "h"
    elif field_name.endswith("_pa"):
        unit = "Pa"
    elif field_name.endswith("_n"):
        unit = "N"

    if field_name in NUMERIC_FIELDS.get(sheet_name, set()):
        data_type = "float"
    elif field_name in {"year", "document_page", "pdf_page", "washing_cycles"}:
        data_type = "integer"
    elif field_name in {"preparation_date", "measurement_date"}:
        data_type = "date"
    else:
        data_type = "text"

    return {
        "sheet_name": sheet_name,
        "field_name": field_name,
        "chinese_name": CHINESE_NAMES.get(field_name, field_name),
        "definition": f"{sheet_name}工作表中的{CHINESE_NAMES.get(field_name, field_name)}字段。",
        "data_type": data_type,
        "unit": unit,
        "allowed_values": allowed or None,
        "required": "yes" if field_name in REQUIRED_FIELDS.get(sheet_name, set()) else "no",
        "missing_value_rule": "blank",
        "example": None,
        "remarks": "80万Da等于0.8 MDa" if field_name == "ha_mw_min_mda" else None,
    }


def build_data_dictionary() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for sheet_name, fields in SHEETS.items():
        if sheet_name == "Validation_Report":
            continue
        for field in fields:
            rows.append(dict_row(sheet_name, field))
    return rows


def add_sheet(ws, columns: list[str], rows: list[dict[str, object]]) -> None:
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        if cell.value in REQUIRED_FIELDS.get(ws.title, set()):
            cell.comment = Comment("required field", "Codex")
    for row in rows:
        ws.append([row.get(col) for col in columns])
    if ws.max_row >= 1 and ws.max_column >= 1:
        ref = f"A1:{chr(64 + ws.max_column)}{ws.max_row}"
        if ws.max_column > 26:
            from openpyxl.utils import get_column_letter
            ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=f"{ws.title}Table", ref=ref)
        table.tableStyleInfo = TABLE_STYLE
        ws.add_table(table)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2" if ws.title not in {"Formulation", "Process"} else "B2"
    style_sheet(ws)
    add_validations(ws, columns)
    if ws.title in {"Formulation", "Process"}:
        ws.freeze_panes = "B2"
    add_duplicate_rules(ws, columns)


def style_sheet(ws) -> None:
    from openpyxl.utils import get_column_letter

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = WRAP
            if cell.row > 1:
                field = ws.cell(1, cell.column).value
                if field in NUMERIC_FIELDS.get(ws.title, set()) and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
                if field in ID_FIELDS.get(ws.title, set()):
                    cell.number_format = "@"
    for col_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 35)


def add_validations(ws, columns: list[str]) -> None:
    for field, values in ENUMS.items():
        if field not in columns or ws.max_row < 2:
            continue
        col_idx = columns.index(field) + 1
        from openpyxl.utils import get_column_letter
        col_letter = get_column_letter(col_idx)
        dv = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
        dv.add(f"{col_letter}2:{col_letter}{max(ws.max_row, 200)}")
        ws.add_data_validation(dv)


def add_duplicate_rules(ws, columns: list[str]) -> None:
    target_fields = [f for f in ["formulation_id", "batch_id", "particle_result_id", "hydrogel_result_id"] if f in columns]
    if ws.max_row < 3:
        return
    from openpyxl.utils import get_column_letter

    for field in target_fields:
        idx = columns.index(field) + 1
        letter = get_column_letter(idx)
        formula = f'COUNTIF(${letter}$2:${letter}${ws.max_row},{letter}2)>1'
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{ws.max_row}",
            FormulaRule(formula=[formula], stopIfTrue=False, fill=SUB_FILL),
        )


def copy_original_data(ws, source_paths: list[Path]) -> None:
    ws["A1"] = "该工作表为原始数据快照，仅用于追溯，不直接用于建模。"
    ws["A1"].fill = SUB_FILL
    ws["A1"].font = Font(bold=True)
    row_cursor = 3
    for path in source_paths:
        ws.cell(row_cursor, 1, f"SOURCE: {path.name}")
        ws.cell(row_cursor, 1).fill = SUB_FILL
        ws.cell(row_cursor, 1).font = Font(bold=True)
        row_cursor += 1
        if path.suffix.lower() == ".xlsx":
            wb = load_workbook(path, read_only=True, data_only=False)
            for sname in wb.sheetnames:
                source_ws = wb[sname]
                ws.cell(row_cursor, 1, f"SHEET: {sname}")
                ws.cell(row_cursor, 1).fill = SUB_FILL
                row_cursor += 1
                for row in source_ws.iter_rows(values_only=True):
                    for col_idx, value in enumerate(row, start=1):
                        ws.cell(row_cursor, col_idx, value)
                    row_cursor += 1
                row_cursor += 2
        else:
            with path.open("r", encoding="utf-8-sig", newline="") as f:
                for row in csv.reader(f):
                    for col_idx, value in enumerate(row, start=1):
                        ws.cell(row_cursor, col_idx, value)
                    row_cursor += 1
                row_cursor += 2
    ws.freeze_panes = "A2"
    style_sheet(ws)


def build_readme_sheet(ws) -> None:
    lines = [
        "Chitosan hydrogel dataset template v2",
        "",
        "Purpose:",
        "- Standardized formulation/process/result workbook for statistics, visualization, and ML.",
        "- Experimental data and literature references are separated.",
        "- Missing values must remain blank.",
        "",
        "Rules:",
        "- Every numeric cell stores number only; units live in headers or Data_Dictionary.",
        "- Do not mix mean±sd in one cell.",
        "- Distinguish independent experimental replicate from technical replicate.",
        "- Hydrogel release rows are one timepoint per row.",
        "",
        "Scientific note:",
        "- 学长论文中的80–150万Da是透明质酸HA的分子量，不是壳聚糖CS的分子量。",
        "",
        "Current source assumption:",
        "- The exact workbook named in the pasted task was not present in the workspace.",
        "- This v2 workbook was built from the current chitosan source files in docs/tables/chitosan/.",
        "",
        "Sheet relations:",
        "- Formulation -> Process via formulation_id.",
        "- Process -> Particle_Result via batch_id.",
        "- Process -> Hydrogel_Result via batch_id.",
        "- Literature_Reference stays separate from experiment tables.",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(i, 1, line)
        ws.cell(i, 1).alignment = WRAP
    ws["A1"].fill = HEADER_FILL
    ws["A1"].font = HEADER_FONT
    ws.column_dimensions["A"].width = 35
    ws.freeze_panes = "A2"


def build_validation_issues(data: dict[str, list[dict[str, object]]], workbook: Workbook) -> list[dict[str, object]]:
    issues: list[dict[str, object]] = []
    counter = 1

    def add(sheet: str, row_number: int | None, field: str, severity: str, problem: str, action: str) -> None:
        nonlocal counter
        issues.append(
            {
                "check_id": f"CHK-{counter:03d}",
                "sheet_name": sheet,
                "row_number": row_number,
                "field_name": field,
                "severity": severity,
                "problem": problem,
                "suggested_action": action,
            }
        )
        counter += 1

    for sheet_name, id_field in [
        ("Formulation", "formulation_id"),
        ("Process", "batch_id"),
        ("Particle_Result", "particle_result_id"),
        ("Hydrogel_Result", "hydrogel_result_id"),
    ]:
        seen = set()
        for idx, row in enumerate(data[sheet_name], start=2):
            value = row.get(id_field)
            if not value:
                add(sheet_name, idx, id_field, "error", "Missing required ID.", "Fill unique ID.")
            elif value in seen:
                add(sheet_name, idx, id_field, "error", "Duplicate ID.", "Deduplicate ID values.")
            seen.add(value)

    formulation_ids = {row["formulation_id"] for row in data["Formulation"] if row["formulation_id"]}
    batch_ids = {row["batch_id"] for row in data["Process"] if row["batch_id"]}
    for idx, row in enumerate(data["Process"], start=2):
        if row.get("formulation_id") not in formulation_ids:
            add("Process", idx, "formulation_id", "error", "Foreign key not found in Formulation.", "Fix formulation_id.")
    for sheet_name in ["Particle_Result", "Hydrogel_Result"]:
        for idx, row in enumerate(data[sheet_name], start=2):
            if row.get("batch_id") not in batch_ids:
                add(sheet_name, idx, "batch_id", "error", "Foreign key not found in Process.", "Fix batch_id.")

    for sheet_name, rows in data.items():
        for idx, row in enumerate(rows, start=2):
            for field in NUMERIC_FIELDS.get(sheet_name, set()):
                value = row.get(field)
                if value is None or value == "":
                    continue
                if not isinstance(value, (int, float)):
                    add(sheet_name, idx, field, "error", "Numeric field contains non-numeric value.", "Convert to pure number or blank.")
                if isinstance(value, str) and re.search(r"[±%A-Za-z]", value):
                    add(sheet_name, idx, field, "error", "Numeric field still contains unit or symbol.", "Strip unit/symbol.")

    for idx, row in enumerate(data["Hydrogel_Result"], start=2):
        pct = row.get("dex_cumulative_release_mean_pct")
        if isinstance(pct, (int, float)) and pct > 100:
            add("Hydrogel_Result", idx, "dex_cumulative_release_mean_pct", "warning", "Percent exceeds 100.", "Review denominator, standard curve, or assay chain.")
    for idx, row in enumerate(data["Particle_Result"], start=2):
        pdi = row.get("pdi_mean")
        if isinstance(pdi, (int, float)) and (pdi < 0 or pdi > 1):
            add("Particle_Result", idx, "pdi_mean", "warning", "PDI outside expected 0-1 range.", "Review particle result.")
        size = row.get("particle_size_mean_nm")
        if isinstance(size, (int, float)) and size <= 0:
            add("Particle_Result", idx, "particle_size_mean_nm", "error", "Particle size <= 0.", "Fix numeric value.")
        for field in ["particle_size_sd_nm", "pdi_sd", "zeta_potential_sd_mv", "dex_recovered_mass_sd_mg", "dex_drug_loading_sd_pct", "dex_encapsulation_efficiency_sd_pct"]:
            val = row.get(field)
            if isinstance(val, (int, float)) and val < 0:
                add("Particle_Result", idx, field, "error", "Standard deviation < 0.", "Fix standard deviation.")

    for idx, row in enumerate(data["Formulation"], start=2):
        lo = row.get("ha_mw_min_mda")
        hi = row.get("ha_mw_max_mda")
        if isinstance(lo, (int, float)) and isinstance(hi, (int, float)) and lo > hi:
            add("Formulation", idx, "ha_mw_min_mda", "error", "HA molecular weight min > max.", "Swap min/max.")
        if row.get("cs_mw_min_kda") or row.get("cs_mw_max_kda"):
            add("Formulation", idx, "cs_mw_min_kda", "warning", "CS MW should remain blank when not reported.", "Clear CS MW fields.")

    # Literature references must not enter experiment tables
    for idx, row in enumerate(data["Particle_Result"], start=2):
        if str(row.get("batch_id", "")).startswith("REF-"):
            add("Particle_Result", idx, "batch_id", "error", "Literature reference mixed into experiment result table.", "Keep literature-only values in Literature_Reference.")

    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    add(ws.title, cell.row, cell.column_letter, "warning", "Formula detected in output workbook.", "Store values only in dataset tables.")

    # Reopen validation
    try:
        load_workbook(OUTPUT_XLSX, read_only=True, data_only=False)
        add("Workbook", 0, "openpyxl", "info", "Workbook reopened successfully with openpyxl.", "No action needed.")
    except Exception as exc:  # pragma: no cover
        add("Workbook", 0, "openpyxl", "error", f"Workbook failed to reopen: {exc}", "Inspect workbook corruption.")

    if not issues:
        add("Workbook", 0, "summary", "info", "No validation issues found.", "No action needed.")
    return issues


def save_validation_report_txt(issues: list[dict[str, object]]) -> None:
    error_count = sum(1 for x in issues if x["severity"] == "error")
    warning_count = sum(1 for x in issues if x["severity"] == "warning")
    info_count = sum(1 for x in issues if x["severity"] == "info")
    lines = [
        f"Validation report generated at {datetime.now().isoformat()}",
        f"errors: {error_count}",
        f"warnings: {warning_count}",
        f"info: {info_count}",
        "",
    ]
    for issue in issues:
        lines.append(
            f"[{issue['severity']}] {issue['sheet_name']} row={issue['row_number']} field={issue['field_name']} :: {issue['problem']} -> {issue['suggested_action']}"
        )
    VALIDATION_TXT.write_text("\n".join(lines), encoding="utf-8")


def save_changelog(backups: list[Path], issue_counts: dict[str, int]) -> None:
    text = f"""# CHANGELOG

## What was changed

- Built a new standardized workbook: `chitosan_hydrogel_dataset_template_v2.xlsx`
- Created sheets: README, Formulation, Process, Particle_Result, Hydrogel_Result, Literature_Reference, Data_Dictionary, Original_Data, Validation_Report
- Preserved source snapshots in `Original_Data`
- Backed up source files into `{BACKUP_DIR}`

## Unit corrections

- Corrected HA molecular-weight semantics from `20-40` / `40-80` to `0.2-0.4 MDa` / `0.4-0.8 MDa`
- Preserved the note that `80-150万 Da` belongs to HA, not CS
- Left `cs_mw_min_kda`, `cs_mw_max_kda`, and `cs_deacetylation_degree_pct` blank because they are not reliably reported

## Field splits

- Separated DEX drug-loading percentages from actual recovered mass fields
- Kept actual recovered mass blank where the source only reported DL%
- Converted feed masses to pure numeric fields: `dex_feed_mass_mg`, `gcv_feed_mass_mg`
- Stored release time series as one row per timepoint in `Hydrogel_Result`

## Left blank because not confirmable

- CS molecular-weight and deacetylation fields
- Most batch/process variables
- Experiment-side particle characterization values not yet reported
- Raw absorbance, blank absorbance, and dilution factors where the source only contains mean Qt

## Literature handling

- Moved thesis reference values into `Literature_Reference`
- Marked reference formulations as `record_type=literature_reference`
- Did not mix literature result values into experiment result sheets

## Potential issues found

- Percent-over-100 release values remain in the imported DEX release series and are flagged in `Validation_Report`
- Group-to-formulation mapping was locked from the orthogonal design table: group1→F2, group2→F3, group3→F5, group4→F6, group5→F1, group6→F4
- The exact workbook name from the pasted task was not present in the workspace; the v2 workbook was built from the current chitosan source files instead

## Backups created

{chr(10).join(f"- {p.name}" for p in backups)}

## Validation summary

- errors: {issue_counts['error']}
- warnings: {issue_counts['warning']}
- info: {issue_counts['info']}
"""
    CHANGELOG_MD.write_text(text, encoding="utf-8")


def main() -> None:
    for path in [SOURCE_WORKBOOK, SOURCE_SAMPLE_XLSX, SOURCE_RELEASE_XLSX, SOURCE_SAMPLE_CSV, SOURCE_RELEASE_CSV]:
        if not path.exists():
            raise FileNotFoundError(f"Missing required source file: {path}")

    backups = backup_sources()
    sample_rows = read_csv(SOURCE_SAMPLE_CSV)
    release_rows = read_csv(SOURCE_RELEASE_CSV)

    formulation_rows, batch_map = build_formulation_rows(sample_rows)
    process_rows = build_process_rows(formulation_rows, batch_map)
    particle_rows = build_particle_rows(formulation_rows, batch_map)
    hydrogel_rows = build_hydrogel_rows(release_rows, batch_map)
    literature_rows = build_literature_rows(sample_rows)
    dictionary_rows = build_data_dictionary()

    data = {
        "Formulation": formulation_rows,
        "Process": process_rows,
        "Particle_Result": particle_rows,
        "Hydrogel_Result": hydrogel_rows,
        "Literature_Reference": literature_rows,
        "Data_Dictionary": dictionary_rows,
    }

    wb = Workbook()
    wb.remove(wb.active)

    ws_readme = wb.create_sheet("README")
    build_readme_sheet(ws_readme)

    for sheet_name in ["Formulation", "Process", "Particle_Result", "Hydrogel_Result", "Literature_Reference", "Data_Dictionary"]:
        ws = wb.create_sheet(sheet_name)
        add_sheet(ws, SHEETS[sheet_name], data[sheet_name])

    ws_original = wb.create_sheet("Original_Data")
    copy_original_data(ws_original, [SOURCE_WORKBOOK, SOURCE_SAMPLE_CSV, SOURCE_RELEASE_CSV])

    ws_validation = wb.create_sheet("Validation_Report")
    # placeholder header first; actual rows after validation
    add_sheet(ws_validation, SHEETS["Validation_Report"], [])

    wb.save(OUTPUT_XLSX)

    reopened = load_workbook(OUTPUT_XLSX)
    issues = build_validation_issues(data, reopened)

    ws_validation = reopened["Validation_Report"]
    if ws_validation.max_row > 1:
        ws_validation.delete_rows(2, ws_validation.max_row - 1)
    for issue in issues:
        ws_validation.append([issue.get(col) for col in SHEETS["Validation_Report"]])
    style_sheet(ws_validation)
    ws_validation.auto_filter.ref = ws_validation.dimensions
    reopened.save(OUTPUT_XLSX)

    save_validation_report_txt(issues)
    counts = {
        "error": sum(1 for x in issues if x["severity"] == "error"),
        "warning": sum(1 for x in issues if x["severity"] == "warning"),
        "info": sum(1 for x in issues if x["severity"] == "info"),
    }
    save_changelog(backups, counts)

    total_rows = sum(len(v) for v in data.values())
    print(f"Output file: {OUTPUT_XLSX}")
    print(f"Backup file: {BACKUP_DIR}")
    print("Sheets created: README, Formulation, Process, Particle_Result, Hydrogel_Result, Literature_Reference, Data_Dictionary, Original_Data, Validation_Report")
    print(f"Rows migrated: {total_rows}")
    print(f"Warnings: {counts['warning']}")
    print(f"Errors: {counts['error']}")
    status = "pass_with_warnings" if counts["error"] == 0 else "fail_with_errors"
    print(f"Validation status: {status}")


if __name__ == "__main__":
    main()
